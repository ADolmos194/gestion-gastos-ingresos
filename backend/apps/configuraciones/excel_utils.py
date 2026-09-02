"""Utilidades para armar los .xlsx de plantilla/exportación (pandas + openpyxl) — genérico
a propósito, ya que todo maestro tipo-planilla del módulo Configuraciones (Categorías, y a
futuro Cuentas/Monedas/etc.) arma sus archivos con el mismo estilo, no solo Categorías.
"""

from io import BytesIO

import pandas as pd
from django.http import HttpResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Mismos colores que ya usa la grilla en pantalla (verde para Activo, gris tachado para
# Cancelado) — para que el Excel no desentone con lo que el usuario ve en la app.
HEADER_FILL = PatternFill(start_color="111827", end_color="111827", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
BODY_FONT = Font(size=10.5)
ZEBRA_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
ACTIVE_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
ACTIVE_FONT = Font(color="166534", bold=True, size=10.5)
CANCELLED_FILL = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
CANCELLED_FONT = Font(color="6B7280", italic=True, size=10.5, strike=True)
_THIN_SIDE = Side(style="thin", color="E5E7EB")
CELL_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)


def style_worksheet(ws: Worksheet, df: pd.DataFrame, active_status_name: str, cancelled_status_name: str) -> None:
    """Le da formato a la hoja que arma pandas (que por defecto sale sin ningún estilo):
    encabezado resaltado y congelado, columnas anchas según el contenido, filas
    intercaladas, autofiltro y, si hay una columna "Estado", coloreada como en la grilla."""
    status_col_index = next((idx for idx, name in enumerate(df.columns, start=1) if name == "Estado"), None)

    for col_idx, column_name in enumerate(df.columns, start=1):
        col_letter = get_column_letter(col_idx)
        header_cell = ws.cell(row=1, column=col_idx)
        header_cell.fill = HEADER_FILL
        header_cell.font = HEADER_FONT
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        header_cell.border = CELL_BORDER

        content_lengths = [len(str(value)) for value in df[column_name]] if len(df) else [0]
        max_len = max(len(str(column_name)), *content_lengths)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

    for row_idx in range(2, len(df) + 2):
        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = CELL_BORDER
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="center")
            if row_idx % 2 == 0:
                cell.fill = ZEBRA_FILL

        if status_col_index:
            status_cell = ws.cell(row=row_idx, column=status_col_index)
            status_cell.alignment = Alignment(horizontal="center", vertical="center")
            if status_cell.value == active_status_name:
                status_cell.fill = ACTIVE_FILL
                status_cell.font = ACTIVE_FONT
            elif status_cell.value == cancelled_status_name:
                status_cell.fill = CANCELLED_FILL
                status_cell.font = CANCELLED_FONT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 22
    ws.sheet_view.showGridLines = False


def dataframe_to_xlsx_response(
    df: pd.DataFrame,
    filename: str,
    sheet_name: str,
    active_status_name: str,
    cancelled_status_name: str,
) -> HttpResponse:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        style_worksheet(writer.sheets[sheet_name], df, active_status_name, cancelled_status_name)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
